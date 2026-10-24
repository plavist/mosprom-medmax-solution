"""Извлечение текста: текстовый слой -> иначе OCR (Tesseract)."""
import io
import os
import subprocess
import tempfile
from pathlib import Path

from . import config

MIN_TEXT_LAYER = 40  # символов на страницу, ниже — считаем страницу сканом


def ocr_image_bytes(data: bytes) -> str:
    """Tesseract по байтам картинки."""
    with tempfile.NamedTemporaryFile(suffix=".png") as f:
        f.write(data)
        f.flush()
        return _tesseract(f.name)


def _tesseract(path: str) -> str:
    r = subprocess.run(
        ["tesseract", path, "-", "-l", config.ocr_lang(), "--psm", "3"],
        capture_output=True, text=True, timeout=180,
    )
    if r.returncode != 0:
        raise RuntimeError(f"tesseract failed: {r.stderr[:300]}")
    return r.stdout.strip()


def _pdf_pages_count(path: str) -> int:
    r = subprocess.run(["pdfinfo", path], capture_output=True, text=True, timeout=60)
    for line in r.stdout.splitlines():
        if line.startswith("Pages:"):
            return int(line.split()[1])
    raise RuntimeError("cannot read page count")


def _pdf_text_layer(path: str, page: int) -> str:
    r = subprocess.run(
        ["pdftotext", "-layout", "-f", str(page), "-l", str(page), path, "-"],
        capture_output=True, text=True, timeout=60,
    )
    return r.stdout.strip()


def _pdf_ocr_page(path: str, page: int) -> str:
    with tempfile.TemporaryDirectory() as tmp:
        prefix = os.path.join(tmp, "p")
        subprocess.run(
            ["pdftoppm", "-r", "200", "-png", "-gray", "-f", str(page), "-l", str(page), path, prefix],
            check=True, capture_output=True, timeout=120,
        )
        imgs = sorted(Path(tmp).glob("p*.png"))
        return _tesseract(str(imgs[0])) if imgs else ""


def extract_pdf(path: str) -> list[tuple[int, str, str]]:
    """[(page, text, method)] где method = 'text' | 'ocr'."""
    out = []
    for page in range(1, _pdf_pages_count(path) + 1):
        text = _pdf_text_layer(path, page)
        method = "text"
        if len(text) < MIN_TEXT_LAYER:
            text, method = _pdf_ocr_page(path, page), "ocr"
        if text.strip():
            out.append((page, text, method))
    return out


def extract_pptx(path: str) -> list[tuple[int, str, str]]:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    out = []
    prs = Presentation(path)
    for i, slide in enumerate(prs.slides, 1):
        parts, method = [], "text"

        def walk(shapes):
            nonlocal method
            for sh in shapes:
                if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
                    walk(sh.shapes)
                elif sh.has_text_frame and sh.text_frame.text.strip():
                    parts.append(sh.text_frame.text.strip())
                elif getattr(sh, "has_table", False) and sh.has_table:
                    for row in sh.table.rows:
                        parts.append(" | ".join(c.text.strip() for c in row.cells))
                elif sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    try:
                        txt = ocr_image_bytes(sh.image.blob)
                    except Exception:
                        txt = ""
                    if len(txt) >= MIN_TEXT_LAYER:
                        parts.append(txt)
                        method = "text+ocr"

        walk(slide.shapes)
        if parts:
            out.append((i, "\n".join(parts), method))
    return out


def extract_docx(path: str) -> list[tuple[int, str, str]]:
    import docx

    d = docx.Document(path)
    lines = [p.text for p in d.paragraphs if p.text.strip()]
    for t in d.tables:
        for row in t.rows:
            lines.append(" | ".join(c.text.strip() for c in row.cells))
    return [(1, "\n".join(lines), "text")] if lines else []


def extract_text_file(path: str) -> list[tuple[int, str, str]]:
    text = Path(path).read_text(encoding="utf-8", errors="replace").strip()
    return [(1, text, "text")] if text else []


def extract_image(path: str) -> list[tuple[int, str, str]]:
    text = _tesseract(path)
    return [(1, text, "ocr")] if text else []


EXTRACTORS = {
    ".pdf": extract_pdf,
    ".pptx": extract_pptx,
    ".docx": extract_docx,
    ".txt": extract_text_file,
    ".md": extract_text_file,
    ".png": extract_image,
    ".jpg": extract_image,
    ".jpeg": extract_image,
}


def supported(filename: str) -> bool:
    return Path(filename).suffix.lower() in EXTRACTORS


def extract(path: str, filename: str) -> list[tuple[int, str, str]]:
    ext = Path(filename).suffix.lower()
    if ext not in EXTRACTORS:
        raise ValueError(f"unsupported file type: {ext}")
    return EXTRACTORS[ext](path)
