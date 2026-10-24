import subprocess
from types import SimpleNamespace

import pytest

from app import extract


def fake_run_factory(pages_text: dict[int, str], ocr_text="OCR TEXT " * 10, tess_rc=0):
    def fake_run(cmd, **kw):
        prog = cmd[0]
        if prog == "pdfinfo":
            return SimpleNamespace(stdout=f"Pages: {len(pages_text)}\n", returncode=0)
        if prog == "pdftotext":
            return SimpleNamespace(stdout=pages_text[int(cmd[cmd.index("-f") + 1])], returncode=0)
        if prog == "pdftoppm":
            open(cmd[-1] + "-1.png", "wb").write(b"x")
            return SimpleNamespace(stdout="", returncode=0)
        if prog == "tesseract":
            return SimpleNamespace(stdout=ocr_text, stderr="boom", returncode=tess_rc)
        raise AssertionError(cmd)
    return fake_run


def test_pdf_text_layer_and_ocr_fallback(monkeypatch):
    long = "Это текстовый слой страницы, достаточно длинный для порога."
    monkeypatch.setattr(subprocess, "run", fake_run_factory({1: long, 2: "", 3: "  "}))
    res = extract.extract_pdf("x.pdf")
    assert [(p, m) for p, _, m in res] == [(1, "text"), (2, "ocr"), (3, "ocr")]
    assert res[0][1] == long


def test_pdf_ocr_empty_page_dropped(monkeypatch):
    monkeypatch.setattr(subprocess, "run", fake_run_factory({1: ""}, ocr_text=" "))
    assert extract.extract_pdf("x.pdf") == []


def test_pdfinfo_failure(monkeypatch):
    monkeypatch.setattr(subprocess, "run", lambda *a, **k: SimpleNamespace(stdout="nothing", returncode=0))
    with pytest.raises(RuntimeError):
        extract.extract_pdf("x.pdf")


def test_tesseract_error(monkeypatch):
    monkeypatch.setattr(subprocess, "run", fake_run_factory({1: ""}, tess_rc=1))
    with pytest.raises(RuntimeError, match="tesseract failed"):
        extract.extract_pdf("x.pdf")


def test_ocr_image_bytes_and_image(monkeypatch, tmp_path):
    monkeypatch.setattr(subprocess, "run", fake_run_factory({}))
    assert extract.ocr_image_bytes(b"png").startswith("OCR TEXT")
    img = tmp_path / "a.png"
    img.write_bytes(b"x")
    assert extract.extract(str(img), "a.png")[0][2] == "ocr"


def test_pptx_text_table_group_and_picture_ocr(monkeypatch, tmp_path):
    from pptx import Presentation
    from pptx.util import Inches
    from PIL import Image

    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Антикоррупционная политика"
    tb = s.shapes.add_table(1, 2, Inches(1), Inches(2), Inches(4), Inches(1)).table
    tb.cell(0, 0).text, tb.cell(0, 1).text = "А", "Б"
    grp = s.shapes.add_group_shape()
    grp.shapes.add_textbox(0, 0, Inches(1), Inches(1)).text_frame.text = "Внутри группы"
    img = tmp_path / "i.png"
    Image.new("RGB", (10, 10), "white").save(img)
    s.shapes.add_picture(str(img), 0, 0)
    s2 = prs.slides.add_slide(prs.slide_layouts[6])  # пустой слайд
    p = tmp_path / "t.pptx"
    prs.save(p)

    monkeypatch.setattr(extract, "ocr_image_bytes", lambda b: "Текст со скана картинки, достаточно длинный текст")
    res = extract.extract_pptx(str(p))
    assert len(res) == 1
    page, text, method = res[0]
    assert "Антикоррупционная политика" in text and "А | Б" in text and "Внутри группы" in text
    assert "скана картинки" in text and method == "text+ocr"


def test_pptx_picture_ocr_failure_ignored(monkeypatch, tmp_path):
    from pptx import Presentation
    from PIL import Image

    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    img = tmp_path / "i.png"
    Image.new("RGB", (10, 10)).save(img)
    s.shapes.add_picture(str(img), 0, 0)
    p = tmp_path / "t.pptx"
    prs.save(p)

    def boom(b):
        raise RuntimeError("no tesseract")
    monkeypatch.setattr(extract, "ocr_image_bytes", boom)
    assert extract.extract_pptx(str(p)) == []


def test_docx(tmp_path):
    import docx

    d = docx.Document()
    d.add_paragraph("Первый абзац")
    t = d.add_table(rows=1, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "x", "y"
    p = tmp_path / "a.docx"
    d.save(p)
    (_, text, _), = extract.extract(str(p), "a.docx")
    assert "Первый абзац" in text and "x | y" in text
    d2 = docx.Document()
    d2.save(tmp_path / "e.docx")
    assert extract.extract_docx(str(tmp_path / "e.docx")) == []


def test_txt_and_unsupported(tmp_path):
    f = tmp_path / "a.txt"
    f.write_text("hello")
    assert extract.extract(str(f), "a.TXT")[0][1] == "hello"
    f.write_text(" ")
    assert extract.extract_text_file(str(f)) == []
    assert extract.supported("x.pdf") and not extract.supported("x.exe")
    with pytest.raises(ValueError):
        extract.extract(str(f), "x.exe")
